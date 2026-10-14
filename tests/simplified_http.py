"""Real HTTP checks on port 8001, using fictional accounts and the UI-test DB only."""
import csv
import hashlib
import http.cookiejar
import io
import json
import os
import re
import uuid
from pathlib import Path
from statistics import median
from time import perf_counter
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener

os.environ['DJANGO_SETTINGS_MODULE'] = 'config.ui_test_settings'
import django
django.setup()

from django.conf import settings
from django.contrib.auth import get_user_model
from openpyxl import Workbook, load_workbook
from coding.models import Assignment, AuditEvent, Code, Codebook, Membership, Project, Round, SourceDocument, SourceRecord, Unit
from coding.services import annotation_data, create_round

if Path(settings.DATABASES['default']['NAME']).resolve() != (settings.RUNTIME_DIR / 'ui-smoke.sqlite3').resolve():
    raise SystemExit('Only the separate UI-test database is allowed.')

cookies = http.cookiejar.CookieJar()
opener = build_opener(HTTPCookieProcessor(cookies))
checks = []
latencies = []


def check(name, condition):
    if not condition:
        raise AssertionError(name)
    checks.append(name)


def request(path, data=None, *, json_body=False, file=None, csrf=True):
    headers, body = {}, None
    if data is not None:
        data = dict(data)
        if csrf:
            token = next(c.value for c in cookies if c.name == 'ui_csrftoken')
            if json_body:
                headers['X-CSRFToken'] = token
            else:
                data['csrfmiddlewaretoken'] = token
        if file is not None:
            boundary = 'fiction-' + uuid.uuid4().hex
            parts = [f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n'.encode()
                     for key, value in data.items()]
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="fiction.xlsx"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode() + file + b'\r\n')
            parts.append(f'--{boundary}--\r\n'.encode())
            body = b''.join(parts)
            headers['Content-Type'] = 'multipart/form-data; boundary=' + boundary
        elif json_body:
            body = json.dumps(data).encode()
            headers['Content-Type'] = 'application/json'
        else:
            body = urlencode(data).encode()
            headers['Content-Type'] = 'application/x-www-form-urlencoded'
    started = perf_counter()
    request_status = None
    try:
        with opener.open(Request('http://127.0.0.1:8001' + path, data=body, headers=headers), timeout=15) as response:
            request_status = response.status
            return response.status, response.read().decode('utf-8-sig'), response.geturl()
    except HTTPError as error:
        request_status = error.code
        return error.code, error.read().decode('utf-8'), error.geturl()
    finally:
        latencies.append({'path': path, 'status': request_status, 'ms': round((perf_counter() - started) * 1000, 2)})


status, _, _ = request('/login/')
check('isolated_server', status == 200 and any(c.name == 'ui_csrftoken' for c in cookies))
suffix = uuid.uuid4().hex
password = uuid.uuid4().hex + '-Fiction!'
owner = get_user_model().objects.create_user('ui_simple_owner_' + suffix, password=password)
coder = get_user_model().objects.create_user('ui_simple_coder_' + suffix, password=password)
project = Project.objects.create(name='简便工作台 · 纯虚构HTTP验收', owner=owner)
for user, role in [(owner, 'owner'), (coder, 'coder')]:
    Membership.objects.create(project=project, user=user, role=role)
book = Codebook.objects.create(project=project, version=1, title='纯虚构规则', policy='multi')
plan = Code.objects.create(book=book, key='PLAN', name='计划', definition='虚构：确定下一步', positive_example='先明确目标。')
checking = Code.objects.create(book=book, key='CHECK', name='检查', definition='虚构：核对结果')
adjustment = Code.objects.create(book=book, key='ADJUST', name='调整', definition='虚构：改变方法')


def login(user):
    status, _, _ = request('/login/', {'username': user.username, 'password': password})
    check('login_' + user.username.split('_')[2], status == 200)


login(owner)
status, page, _ = request(f'/p/{project.pk}/')
check('overview_has_direct_material_import', status == 200 and 'id="project-material-file"' in page
      and f'action="/p/{project.pk}/materials/"' in page)
workbook = Workbook()
workbook.active.title = '说明'
sheet = workbook.create_sheet('对话')
sheet.append(['Message Content', 'conversation_id', 'speaker'])
sheet.append(['先😊讨论，再检查', 'A', '虚构甲'])
sheet.append(['明确目标。', 'A', '虚构乙'])
output = io.BytesIO()
workbook.save(output)
workbook.close()
raw = output.getvalue()
materials_url = f'/p/{project.pk}/materials/'
status, page, _ = request(materials_url, {}, file=raw)
check('multi_sheet_empty_cover', status == 200 and '查看这张表' in page and '对话' in page)
draft = project.uploaddraft_set.get(user=owner)
status, page, _ = request(materials_url, {'draft_id': draft.pk, 'sheet_name': '对话', 'action': 'preview'})
check('switch_sheet_without_import', status == 200 and 'Message Content' in page and not project.documents.exists())
confirm = {'draft_id': draft.pk, 'sheet_name': '对话', 'text_column': 'Message Content', 'context_column': 'conversation_id'}
status, _, _ = request(materials_url, confirm)
doc = project.documents.get()
check('import_preserves_bytes_and_sheet', status == 200 and bytes(doc.raw) == raw and doc.sheet_name == '对话'
      and doc.sha256 == hashlib.sha256(raw).hexdigest() and doc.records.count() == 2)
request(materials_url, confirm)
check('duplicate_confirmation_does_not_duplicate', project.documents.count() == 1)

login(coder)
start_url = f'/p/{project.pk}/quick-start/'
start = {'book_id': book.pk, 'token': str(uuid.uuid4())}
request(start_url, start)
r = Round.objects.get(project=project, kind='quick')
status, page, url = request(f'/p/{project.pk}/code/')
check('continue_directly_with_real_rules', status == 200 and f'/r/{r.pk}/' in url and 'data-code-picker' in page
      and '先明确目标。' in page and '确认并下一条' in page)
a = r.assignments.order_by('id').first()
split_url = f'/api/assignments/{a.pk}/split/'
split = {'revision': 0, 'offset': 2, 'token': str(uuid.uuid4())}
status, _, _ = request(split_url, split, json_body=True, csrf=False)
check('split_requires_csrf', status == 403 and not a.segments.exists())
status, body, _ = request(split_url, split, json_body=True)
check('personal_split_success', status == 200 and json.loads(body)['next_url'].startswith(f'/r/{r.pk}/?a='))
request(split_url, split, json_body=True)
children = list(a.segments.select_related('unit').order_by('unit__start'))
check('split_replay_no_duplicate', len(children) == 2 and AuditEvent.objects.filter(project=project, action='annotation.split').count() == 1)
check('split_unicode_and_original_pool', children[0].unit.text == '先😊' and ''.join(c.unit.text for c in children) == a.unit.text
      and Unit.objects.filter(record__document=doc, active=True).count() == 2)
status, body, _ = request(f'/api/rounds/{r.pk}/progress/')
check('live_progress_without_original_parent', status == 200 and json.loads(body)['my_total'] == 3)
status, _, _ = request(f'/r/{r.pk}/?a={a.pk}')
check('old_link_reopens_live_work', status == 200)
save = {'revision': 0, 'primary': plan.pk, 'secondary': [], 'uncertain': False, 'no_code': False,
        'note': '纯虚构：根据计划定义选择', 'token': str(uuid.uuid4()), 'action': 'submit'}
status, body, _ = request(f'/api/assignments/{children[0].pk}/save/', save, json_body=True)
check('submit_acknowledges_next_segment', status == 200 and json.loads(body)['next_url'] == f'/r/{r.pk}/?a={children[1].pk}')
status, page, _ = request(f'/p/{project.pk}/history/?' + urlencode({'q': '纯虚构'}))
check('searchable_history_after_coding', status == 200 and '纯虚构：根据计划定义选择' in page)
status, csv_text, _ = request(f'/r/{r.pk}/export/mine/')
check('export_tracks_parent_without_double_count', status == 200 and 'parent_assignment_id' in csv_text
      and len(csv_text.strip().splitlines()) == 4)
status, _, _ = request(f'/api/assignments/{a.pk}/save/', {**save, 'token': split['token']}, json_body=True)
check('old_parent_cannot_be_saved', status == 409)

saved_parent = r.assignments.filter(superseded=False, parent__isnull=True).get()
save_url = f'/api/assignments/{saved_parent.pk}/save/'
paired = {**save, 'primary': plan.pk, 'secondary': [checking.pk], 'action': 'save',
          'note': '纯虚构：主次代码一起保存', 'token': str(uuid.uuid4())}
status, body, _ = request(save_url, paired, json_body=True)
check('primary_and_single_secondary_saved_together', status == 200
      and json.loads(body)['annotation']['primary'] == plan.pk
      and json.loads(body)['annotation']['secondary'] == [checking.pk])
paired = {**paired, 'revision': 1, 'primary': checking.pk, 'secondary': [plan.pk], 'token': str(uuid.uuid4())}
status, body, _ = request(save_url, paired, json_body=True)
check('primary_and_single_secondary_can_be_changed', status == 200
      and json.loads(body)['annotation']['primary'] == checking.pk
      and json.loads(body)['annotation']['secondary'] == [plan.pk])
status, page, _ = request(f'/r/{r.pk}/?a={saved_parent.pk}')
embedded = re.search(r'<script id="annotation-data" type="application/json">(.*?)</script>', page, re.S)
reloaded = json.loads(embedded.group(1)) if embedded else {}
check('reloaded_page_keeps_primary_and_one_radio_secondary', status == 200
      and reloaded.get('primary') == checking.pk and reloaded.get('secondary') == [plan.pk]
      and 'type="radio" name="secondary"' in page and 'type="checkbox" name="secondary"' not in page)
saved_parent.refresh_from_db()
old_labels = annotation_data(saved_parent)
status, _, _ = request(save_url, {**paired, 'revision': 2, 'secondary': [plan.pk, adjustment.pk],
                               'token': str(uuid.uuid4())}, json_body=True)
saved_parent.refresh_from_db()
check('multiple_secondaries_rejected_without_overwrite', status == 400 and annotation_data(saved_parent) == old_labels)
selected_split = {'revision': 2, 'start': 1, 'end': 3, 'token': str(uuid.uuid4())}
status, body, _ = request(f'/api/assignments/{saved_parent.pk}/split/', selected_split, json_body=True)
pieces = list(saved_parent.segments.select_related('unit').order_by('unit__start'))
check('saved_middle_selection_has_three_lossless_pieces', status == 200 and len(pieces) == 3
      and [(piece.unit.start, piece.unit.end) for piece in pieces] == [(0, 1), (1, 3), (3, len(saved_parent.unit.text))]
      and ''.join(piece.unit.text for piece in pieces) == saved_parent.unit.text
      and json.loads(body)['next_url'] == f'/r/{r.pk}/?a={pieces[1].pk}')
saved_parent.refresh_from_db()
check('split_parent_keeps_old_primary_secondary_and_note', saved_parent.superseded
      and saved_parent.primary_id == old_labels['primary'] and saved_parent.secondary == old_labels['secondary']
      and saved_parent.note == old_labels['note'])
status, page, _ = request(f'/p/{project.pk}/history/{saved_parent.pk}/')
check('split_parent_history_is_read_only_with_old_labels', status == 200 and '原记录已拆分' in page
      and 'CHECK · 检查' in page and 'PLAN · 计划' in page and old_labels['note'] in page
      and f'href="/r/{r.pk}/?a={saved_parent.pk}"' not in page)
status, page, _ = request(f'/r/{r.pk}/?a={pieces[1].pk}')
check('selected_piece_opens_blank_coding_with_parent_history_link', status == 200
      and f'href="/p/{project.pk}/history/{saved_parent.pk}/"' in page
      and 'data-code-picker' in page and all(not piece.primary_id and not piece.secondary for piece in pieces))
status, _, _ = request(f'/api/assignments/{pieces[1].pk}/save/', {**save, 'action': 'save',
    'secondary': [checking.pk], 'note': '纯虚构：选中片段单独编码', 'token': str(uuid.uuid4())}, json_body=True)
check('selected_piece_records_its_own_primary_and_secondary', status == 200)
status, csv_text, _ = request(f'/r/{r.pk}/export/mine/')
rows = list(csv.DictReader(io.StringIO(csv_text)))
exported = sorted([row for row in rows if row['parent_assignment_id'] == str(saved_parent.pk)],
                  key=lambda row: int(row['span_start']))
check('export_contains_exact_selected_spans_without_parent_double_count', status == 200 and len(exported) == 3
      and ''.join(row['unit_text'] for row in exported) == saved_parent.unit.text
      and exported[1]['span_start'] == '1' and exported[1]['span_end'] == '3'
      and exported[1]['primary_code'] == 'PLAN' and exported[1]['secondary_code'] == 'CHECK'
      and not any(row['assignment_id'] == str(saved_parent.pk) for row in rows))

# A new, fictional fixture in this script's own isolated project exercises many cuts.
metadata = {'created_at': '2026-10-14T11:05:06.123+08:00', 'condition': 'multi_agent',
            'group_id': 'G007', 'sender_id': 'S042', 'message_id': 'M0008', 'seq_num': '0008'}
multi_text = '先😊计划，再检查，然后调整，最后总结。'
multi_csv = io.StringIO(newline='')
writer = csv.writer(multi_csv)
writer.writerow(['text', *metadata])
writer.writerow([multi_text, *metadata.values()])
multi_raw = multi_csv.getvalue().encode('utf-8-sig')
multi_doc = SourceDocument.objects.create(project=project, name='fictional-multisplit.csv',
    raw=multi_raw, sha256=hashlib.sha256(multi_raw).hexdigest(), text_column='text')
multi_record = SourceRecord.objects.create(document=multi_doc, position=1, text=multi_text,
    metadata=metadata, context_key='G007')
multi_unit = Unit.objects.create(record=multi_record, text=multi_text, end=len(multi_text))
multi_parent = Assignment.objects.create(round=r, unit=multi_unit, coder=coder)
original_record = SourceRecord.objects.values().get(pk=multi_record.pk)
multi_split = {'cuts': [4, 8, 12], 'start': 8, 'revision': 0, 'token': str(uuid.uuid4())}
multi_url = f'/api/assignments/{multi_parent.pk}/split/'
status, body, _ = request(multi_url, multi_split, json_body=True)
multi_parts = list(multi_parent.segments.select_related('unit').order_by('unit__start'))
check('many_cuts_create_four_ordered_complete_units', status == 200 and len(multi_parts) == 4
      and ''.join(piece.unit.text for piece in multi_parts) == multi_text
      and json.loads(body)['next_url'] == f'/r/{r.pk}/?a={multi_parts[2].pk}')
status, _, _ = request(multi_url, multi_split, json_body=True)
check('many_cuts_retry_does_not_duplicate', status == 200 and multi_parent.segments.count() == 4)
status, _, _ = request(multi_url, {**multi_split, 'cuts': [4, 9, 12]}, json_body=True)
check('many_cuts_changed_retry_is_conflict', status == 409 and multi_parent.segments.count() == 4)
summary_code = Code.objects.create(book=book, key='SUMMARY', name='总结', definition='虚构：整理结论')
for piece, code in zip(multi_parts, [plan, checking, adjustment, summary_code]):
    status, _, _ = request(f'/api/assignments/{piece.pk}/save/', {**save, 'primary': code.pk,
        'secondary': [], 'token': str(uuid.uuid4())}, json_body=True)
    check('many_cuts_independent_code_' + code.key, status == 200)
check('many_cuts_source_metadata_and_file_unchanged',
      SourceRecord.objects.values().get(pk=multi_record.pk) == original_record
      and bytes(SourceDocument.objects.get(pk=multi_doc.pk).raw) == multi_raw)
for file_format in ['csv', 'json', 'xlsx']:
    export_path = f'/r/{r.pk}/export/mine/?format={file_format}'
    if file_format == 'xlsx':
        with opener.open(Request('http://127.0.0.1:8001' + export_path), timeout=15) as response:
            export_status, export_raw = response.status, response.read()
        workbook = load_workbook(io.BytesIO(export_raw), read_only=True)
        try:
            values = list(workbook.active.values)
            multi_rows = [dict(zip(values[0], row)) for row in values[1:]]
        finally:
            workbook.close()
    else:
        export_status, export_text, _ = request(export_path)
        multi_rows = (json.loads(export_text)['rows'] if file_format == 'json'
                      else list(csv.DictReader(io.StringIO(export_text, newline=''))))
    selected_rows = [row for row in multi_rows if str(row['document_id']) == str(multi_doc.pk)]
    check('many_cuts_' + file_format + '_codes_and_source_metadata_preserved', export_status == 200
          and len(selected_rows) == 4
          and [row['primary_code'] for row in selected_rows] == ['PLAN', 'CHECK', 'ADJUST', 'SUMMARY']
          and ''.join(row['unit_text'] for row in selected_rows) == multi_text
          and all(json.loads(row['metadata_json']) == metadata for row in selected_rows)
          and all(row['meta_' + key] == value for row in selected_rows for key, value in metadata.items())
          and not any(str(row['assignment_id']) == str(multi_parent.pk) for row in multi_rows))

team = create_round(project, owner, book.pk, '纯虚构多人任务', 'pilot', [owner.pk, coder.pk], 0, 2026, 'all', 25)
team_a = team.assignments.get(coder=coder, unit=a.unit)
snapshot = list(team.assignments.values())
status, body, _ = request(f'/api/assignments/{team_a.pk}/split-suggestion/', {**split, 'reason': '虚构：两个行为需分开'}, json_body=True)
check('shared_suggestion_preserves_all_answers', status == 200 and list(team.assignments.values()) == snapshot)
other_snapshot = list(team.assignments.exclude(coder=coder).values())
status, _, _ = request(f'/api/assignments/{team_a.pk}/split/', split, json_body=True)
team_a.refresh_from_db()
check('shared_round_splits_only_own_task', status == 200 and team_a.superseded and team_a.segments.count() == 2
      and list(team.assignments.exclude(coder=coder).values()) == other_snapshot)
login(owner)
status, page, _ = request(f'/p/{project.pk}/units/')
check('owner_sees_suggestion', status == 200 and '虚构：两个行为需分开' in page)
status, _, _ = request(split_url, split, json_body=True)
check('cannot_split_another_member_task', status == 404)

login(coder)
for entry in r.assignments.filter(superseded=False, status='draft'):
    status, _, _ = request(f'/api/assignments/{entry.pk}/save/', {**save, 'revision': entry.revision,
                           'token': str(uuid.uuid4())}, json_body=True)
    check('submit_remaining_' + str(entry.pk), status == 200)
r.refresh_from_db()
status, _, _ = request(f'/p/{project.pk}/quick/{r.pk}/finish/', {'revision': r.revision})
r.refresh_from_db()
check('personal_archive_ignores_replaced_parent', status == 200 and r.state == 'closed')

save_times = [entry['ms'] for entry in latencies if entry['status'] == 200 and entry['path'].endswith('/save/')]
print(json.dumps({'passed': len(checks), 'checks': checks,
                  'successful_save_http_ms': {'count': len(save_times), 'median': round(median(save_times), 2),
                                              'max': max(save_times)}}, ensure_ascii=False))
