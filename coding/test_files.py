"""Synthetic import/export regression tests; never use the user's research files."""
import csv
import io
import json
import subprocess
import zipfile
from datetime import timedelta
from unittest.mock import patch
from urllib.parse import urlparse, parse_qs
from xml.etree import ElementTree as ET

from django.db.models import F
from django.test import TestCase, Client, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .tests import fixture, Helpers, TEST_SETTINGS
from .models import Codebook, Code, CodebookImportDraft, AuditEvent, Membership
from .codebook_files import FIELDS, parse_codebook, normalized_codes, suggest_mapping, MAX_BYTES
from .exporting import ExportTable, book_table, export_response, docx_bytes, annotation_table
from .services import create_round, project_archive


CSV = '代码标识,代码名称,定义,上级代码标识,正例\nCHILD,子代码,具体行为,ROOT,虚构例子\nROOT,上级代码,整体行为,,\n'.encode('utf-8')
WORD_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def minimal_docx(xml):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', xml)
    return output.getvalue()


def workbook_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def text_pdf(text='key: PLAN\nname: Plan\ndefinition: Choose goals', pages=1, password=None):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/Font'),
               NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')}))
    for index in range(pages):
        page = writer.add_blank_page(width=600, height=800)
        if text:
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            escaped = [line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)') for line in text.split('\n')]
            stream.set_data(('BT /F1 12 Tf 20 TL 50 750 Td ' + ' T* '.join(f'({line}) Tj' for line in escaped) + ' ET').encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class CodebookParserTests(TestCase):
    def parse(self, name, raw, sheet=''):
        headers, rows, _, _, _ = parse_codebook(name, raw, sheet)
        return normalized_codes(rows, suggest_mapping(headers))

    def test_csv_chinese_headers_and_child_before_parent(self):
        rows = self.parse('rules.csv', CSV)
        self.assertEqual(rows[0]['parent_key'], 'ROOT')
        self.assertEqual(rows[1]['name'], '上级代码')

    def test_gb18030_and_tsv(self):
        raw = 'key\tname\tdefinition\n001\t计划\t确定目标\n'.encode('gb18030')
        self.assertEqual(self.parse('rules.tsv', raw)[0]['key'], '001')

    def test_multisheet_auto_select_and_explicit_switch(self):
        raw = workbook_bytes([('说明', [('文件说明',)]), ('编码', [('key', 'name', 'definition'), ('001', '计划', '确定目标')]),
                              ('另一个编码表', [('key', 'name', 'definition'), ('B', '检查', '检查进度')])])
        headers, rows, sheets, selected, _ = parse_codebook('book.xlsx', raw)
        self.assertEqual(selected, '编码')
        self.assertEqual(len(sheets), 3)
        self.assertEqual(rows[0]['key'], '001')
        self.assertEqual(self.parse('book.xlsx', raw, '另一个编码表')[0]['key'], 'B')
        with self.assertRaisesMessage(ValueError, '请选择'):
            self.parse('book.xlsx', raw, '不存在')

    def test_excel_formula_is_text_not_computed(self):
        raw = workbook_bytes([('codes', [('key', 'name', 'definition'), ('A', '测试', '=HYPERLINK("bad","text")')])])
        self.assertTrue(self.parse('book.xlsx', raw)[0]['definition'].startswith('=HYPERLINK'))

    def test_json_portable_and_old_archive_parent_ids(self):
        raw = json.dumps({'codes': [{'id': 7, 'key': 'ROOT', 'name': '上级', 'definition': '总体', 'parent': None},
                                   {'id': 8, 'key': 'CHILD', 'name': '子级', 'definition': '具体', 'parent': 7}]}).encode()
        self.assertEqual(self.parse('old.json', raw)[1]['parent_key'], 'ROOT')

    def test_duplicate_json_fields_are_not_silently_replaced(self):
        with self.assertRaisesMessage(ValueError, '重复字段'):
            self.parse('rules.json', b'[{"key":"A","key":"B"}]')

    def test_text_blocks_multiline_definition_and_aliases(self):
        raw = '代码标识：A\n代码名称：计划\n定义：第一行\n第二行\n正例：虚构例子\n\n代码标识：B\n代码名称：检查\n定义：检查进度'.encode()
        rows = self.parse('rules.txt', raw)
        self.assertEqual(rows[0]['definition'], '第一行\n第二行')
        self.assertEqual(len(rows), 2)

    def test_markdown_table(self):
        raw = b'|key|name|definition|\n|---|---|---|\n|A|Plan|Choose goals|\n'
        self.assertEqual(self.parse('rules.md', raw)[0]['name'], 'Plan')

    def test_docx_table(self):
        xml = f'<w:document xmlns:w="{WORD_NS}"><w:body><w:tbl>' + ''.join(
            '<w:tr>' + ''.join(f'<w:tc><w:p><w:r><w:t>{value}</w:t></w:r></w:p></w:tc>' for value in row) + '</w:tr>'
            for row in [('代码标识', '代码名称', '定义'), ('A', '计划', '确定目标')]) + '</w:tbl></w:body></w:document>'
        self.assertEqual(self.parse('rules.docx', minimal_docx(xml))[0]['key'], 'A')

    def test_exported_word_can_be_reimported_without_losing_line_breaks(self):
        raw = docx_bytes(['key', 'name', 'definition', 'parent_key'],
                         [['A', '计划', '第一行\n第二行', ''], ['B', '检查', '检查进度', 'A']], '虚构规则')
        rows = self.parse('rules.docx', raw)
        self.assertEqual(rows[0]['definition'], '第一行\n第二行')
        self.assertEqual(rows[1]['parent_key'], 'A')

    def test_plain_paper_is_not_converted_into_guessed_codes(self):
        with self.assertRaisesMessage(ValueError, '未识别'):
            self.parse('paper.txt', '这是一段论文正文，没有编码规则字段。'.encode())

    def test_docx_external_entities_rejected(self):
        xml = f'<!DOCTYPE document [<!ENTITY x SYSTEM "file:///secret">]><w:document xmlns:w="{WORD_NS}"><w:body><w:p>&x;</w:p></w:body></w:document>'
        with self.assertRaises(ValueError):
            self.parse('rules.docx', minimal_docx(xml))

    def test_macro_zip_and_overlarge_upload_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('word/vbaProject.bin', b'no')
        with self.assertRaisesMessage(ValueError, '宏'):
            self.parse('rules.docx', output.getvalue())
        with self.assertRaisesMessage(ValueError, '8MB'):
            self.parse('rules.csv', b'x' * (MAX_BYTES + 1))

    def test_text_pdf_actual_worker(self):
        self.assertEqual(self.parse('rules.pdf', text_pdf())[0]['key'], 'PLAN')

    def test_scanned_encrypted_and_too_many_pdf_pages(self):
        for raw, phrase in [(text_pdf(text=''), '没有可读取文字'), (text_pdf(password='test-only'), '加密'),
                            (text_pdf(pages=201), '最多200页')]:
            with self.subTest(phrase=phrase), self.assertRaisesMessage(ValueError, phrase):
                self.parse('rules.pdf', raw)

    def test_pdf_timeout_is_a_useful_error(self):
        with patch('coding.codebook_files.subprocess.run', side_effect=subprocess.TimeoutExpired('pdf', 15)):
            with self.assertRaisesMessage(ValueError, '超时'):
                self.parse('rules.pdf', b'%PDF-1.4')

    def test_duplicate_headers_required_fields_and_mapping_collisions(self):
        for raw in [b'key,key,name,definition\nA,B,Plan,Goal', b'key,name,definition\nA,Plan,']:
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                self.parse('rules.csv', raw)
        with self.assertRaisesMessage(ValueError, '多个编码字段'):
            normalized_codes([{'a': 'hello'}], {'key': 'a', 'name': 'a', 'definition': 'a'})

    def test_missing_parent_cycles_long_key_and_invalid_color(self):
        for row in [{'key': 'A', 'name': '计划', 'definition': '目标', 'parent_key': 'MISSING'},
                    {'key': 'A', 'name': '计划', 'definition': '目标', 'parent_key': 'A'},
                    {'key': 'A' * 41, 'name': '计划', 'definition': '目标'},
                    {'key': 'A', 'name': '计划', 'definition': '目标', 'color': 'red'}]:
            with self.subTest(row=row), self.assertRaises(ValueError):
                normalized_codes([row], {f: f for f, _, _ in FIELDS})

    def test_deep_hierarchy_avoids_recursion(self):
        rows = [{'key': str(i), 'name': str(i), 'definition': '层级', 'parent_key': str(i - 1) if i else ''} for i in range(1000)]
        self.assertEqual(len(normalized_codes(rows, {f: f for f, _, _ in FIELDS})), 1000)

    def test_over_1000_rows_and_rows_with_only_unmapped_data_not_skipped(self):
        with self.assertRaisesMessage(ValueError, '1000'):
            self.parse('rules.csv', b'key,name,definition\n' + b'A,Plan,Goal\n' * 1001)
        with self.assertRaisesMessage(ValueError, '所选编码字段均为空'):
            normalized_codes([{'notes': 'not a code'}], {'key': 'key', 'name': 'name', 'definition': 'definition'})


@override_settings(**TEST_SETTINGS)
class CodebookImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        for key, value in fixture().items():
            setattr(cls, key, value)

    def setUp(self):
        self.client.force_login(self.owner)
        self.url = f'/p/{self.project.id}/books/{self.book.id}/import/'

    def upload(self, raw=CSV, filename='rules.csv', book=None):
        url = f'/p/{self.project.id}/books/{(book or self.book).id}/import/'
        response = self.client.post(url, {'action': 'upload', 'file': SimpleUploadedFile(filename, raw)})
        self.assertEqual(response.status_code, 302, response.content)
        draft_id = parse_qs(urlparse(response.url).query)['draft'][0]
        return CodebookImportDraft.objects.get(pk=draft_id)

    def preview(self, draft, mode='new', **extra):
        headers, _, _, _, _ = parse_codebook(draft.filename, bytes(draft.raw))
        data = {'action': 'preview', 'draft_id': str(draft.id), 'mode': mode, 'title': '导入后的虚构规则',
                'policy': 'multi', 'change_note': '纯虚构导入验收'}
        data.update({'map_' + k: v for k, v in suggest_mapping(headers).items()})
        data.update(extra)
        return self.client.post(f'/p/{self.project.id}/books/{draft.book_id}/import/', data)

    def confirm(self, draft, **extra):
        return self.client.post(f'/p/{self.project.id}/books/{draft.book_id}/import/',
                                {'action': 'confirm', 'draft_id': str(draft.id), 'confirmed': 'yes', **extra})

    def test_complete_preview_confirm_and_immediate_new_round_use(self):
        count = Code.objects.count()
        draft = self.upload()
        response = self.client.get(self.url + '?draft=' + str(draft.pk))
        self.assertContains(response, '对应字段')
        self.assertEqual(Code.objects.count(), count)
        response = self.preview(draft)
        self.assertEqual(response.status_code, 302)
        response = self.client.get(response.url)
        self.assertContains(response, '确认导入结果')
        self.assertContains(response, '具体行为')
        self.assertEqual(Code.objects.count(), count)
        self.assertEqual(self.confirm(draft).status_code, 302)
        draft.refresh_from_db()
        new_book = draft.applied_book
        self.assertEqual(new_book.codes.get(key='CHILD').parent.key, 'ROOT')
        self.assertFalse(new_book.frozen)
        self.assertEqual(self.book.codes.count(), 2)
        self.assertEqual(self.r.book_id, self.book.id)
        r = create_round(self.project, self.owner, new_book.id, '使用导入规则', 'training', [self.owner.id], 1, 9, 'all', 25)
        response = self.client.get(f'/r/{r.id}/')
        self.assertContains(response, '子代码')
        self.assertEqual(AuditEvent.objects.filter(action='codebook.import', project=self.project).count(), 1)

    def test_confirmation_requires_explicit_check_and_is_replay_safe(self):
        draft = self.upload()
        self.preview(draft)
        response = self.confirm(draft, confirmed='no')
        self.assertContains(response, '请勾选')
        draft.refresh_from_db()
        self.assertIsNone(draft.applied_book_id)
        self.confirm(draft)
        count = Code.objects.count()
        self.confirm(draft)
        self.assertEqual(Code.objects.count(), count)
        self.assertEqual(AuditEvent.objects.filter(action='codebook.import').count(), 1)

    def test_append_parent_may_refer_to_existing_code_and_policy_stays(self):
        book = Codebook.objects.create(project=self.project, version=2, title='草稿', policy='single')
        Code.objects.create(book=book, key='ROOT', name='现有上级', definition='现有定义')
        draft = self.upload(b'key,name,definition,parent_key\nCHILD,Child,Specific,ROOT\n', book=book)
        self.assertEqual(self.preview(draft, mode='append').status_code, 302)
        self.confirm(draft)
        book.refresh_from_db()
        self.assertEqual(book.policy, 'single')
        self.assertEqual(book.codes.count(), 2)
        self.assertEqual(book.codes.get(key='ROOT').definition, '现有定义')

    def test_append_duplicate_never_overwrites_or_partially_imports(self):
        book = Codebook.objects.create(project=self.project, version=2)
        Code.objects.create(book=book, key='ROOT', name='现有', definition='保留')
        draft = self.upload(book=book)
        response = self.preview(draft, mode='append')
        self.assertContains(response, '不会覆盖旧代码')
        self.assertEqual(book.codes.count(), 1)
        draft.refresh_from_db()
        self.assertEqual(draft.preview, {})

    def test_frozen_book_does_not_offer_append_and_rejects_forged_append(self):
        draft = self.upload()
        response = self.client.get(self.url + '?draft=' + str(draft.pk))
        self.assertNotContains(response, '<option value="append">')
        response = self.preview(draft, mode='append')
        self.assertEqual(response.status_code, 200)
        draft.refresh_from_db()
        self.assertEqual(draft.preview, {})

    def test_stale_preview_does_not_change_book(self):
        draft = self.upload()
        self.preview(draft)
        Codebook.objects.filter(pk=self.book.id).update(revision=F('revision') + 1)
        response = self.confirm(draft)
        self.assertContains(response, '预览后编码本已变化')
        draft.refresh_from_db()
        self.assertIsNone(draft.applied_book_id)
        self.assertEqual(self.project.codebooks.count(), 1)

    def test_expired_preview_and_confirm_without_preview(self):
        draft = self.upload()
        response = self.confirm(draft)
        self.assertContains(response, '预览后编码本已变化')
        CodebookImportDraft.objects.filter(pk=draft.pk).update(created_at=timezone.now() - timedelta(hours=25))
        response = self.confirm(draft)
        self.assertContains(response, '超过24小时')
        self.assertEqual(self.project.codebooks.count(), 1)

    def test_coder_viewer_and_outsider_cannot_import_and_other_owner_cannot_read_draft(self):
        draft = self.upload()
        for user in (self.coder, self.viewer, self.outsider):
            self.client.force_login(user)
            response = self.client.get(self.url + '?draft=' + str(draft.pk))
            self.assertIn(response.status_code, (403, 404))
        Membership.objects.create(project=self.project, user=self.outsider, role='reviewer')
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(self.url + '?draft=' + str(draft.pk)).status_code, 404)

    def test_foreign_book_id_and_malformed_draft_do_not_leak(self):
        response = self.client.get(f'/p/{self.project.id}/books/{self.foreign_code.book_id}/import/')
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.client.get(self.url + '?draft=not-a-uuid').status_code, 404)

    def test_templates_can_all_be_imported(self):
        for extension in ('xlsx', 'csv', 'docx', 'txt', 'json'):
            response = self.client.get(f'/p/{self.project.id}/books/{self.book.id}/template/{extension}/')
            self.assertEqual(response.status_code, 200)
            headers, rows, _, _, _ = parse_codebook('template.' + extension, response.content)
            self.assertEqual(normalized_codes(rows, suggest_mapping(headers))[0]['key'], 'EXAMPLE')

    def test_archive_keeps_import_original_and_mapping(self):
        draft = self.upload()
        self.preview(draft)
        self.confirm(draft)
        output = project_archive(self.project)
        with zipfile.ZipFile(io.BytesIO(output)) as archive:
            self.assertEqual(archive.read(f'02_codebook/import_{draft.pk}.csv'), CSV)
            manifest = json.loads(archive.read(f'02_codebook/import_{draft.pk}_manifest.json'))
            self.assertEqual(manifest['preview']['mapping']['key'], '代码标识')

    def test_csrf_protects_file_upload(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.owner)
        response = client.post(self.url, {'action': 'upload', 'file': SimpleUploadedFile('rules.csv', CSV)})
        self.assertEqual(response.status_code, 403)


@override_settings(**TEST_SETTINGS)
class ExportTests(Helpers, TestCase):
    @classmethod
    def setUpTestData(cls):
        for key, value in fixture().items():
            setattr(cls, key, value)

    def setUp(self):
        self.client.force_login(self.owner)
        self.url = f'/p/{self.project.id}/exports/'

    def data(self, table, target='codebook', format='xlsx', fields=None):
        fields = fields or table.headers
        data = {'target': target, 'book': str(self.book.id), 'round': str(self.r.id), 'format': format,
                'filename': '虚构导出测试', 'encoding': 'utf-8-sig', 'delimiter': ',', 'fields': fields}
        for index, key in enumerate(table.headers):
            data[f'label_{index}'] = key
            data[f'order_{index}'] = str(index + 1)
        return data

    def test_center_ui_lists_formats_and_custom_fields(self):
        response = self.client.get(self.url)
        self.assertContains(response, '导出中心')
        self.assertContains(response, 'Word文档')
        self.assertContains(response, '导出列名')
        self.assertContains(response, 'id="page-back"')

    def test_every_book_format_and_roundtrip(self):
        table = book_table(self.book)
        for extension in ('xlsx', 'csv', 'tsv', 'json', 'docx'):
            response = self.client.post(self.url, self.data(table, format=extension))
            self.assertEqual(response.status_code, 200)
            self.assertIn('attachment;', response.get('Content-Disposition', ''))
            self.assertEqual(response.get('Cache-Control'), 'no-store')
            headers, rows, _, _, _ = parse_codebook('export.' + extension, response.content)
            codes = normalized_codes(rows, suggest_mapping(headers))
            self.assertEqual({c['key'] for c in codes}, {'PLAN', 'CHECK'})

    def test_select_rename_reorder_and_custom_separator_encoding(self):
        data = self.data(book_table(self.book), format='csv', fields=['key', 'name'])
        data.update(label_0='代码编号', label_1='中文名称', order_0='2', order_1='1', delimiter=';', encoding='gb18030')
        response = self.client.post(self.url, data)
        rows = list(csv.reader(io.StringIO(response.content.decode('gb18030')), delimiter=';'))
        self.assertEqual(rows[0], ['中文名称', '代码编号'])
        self.assertEqual(rows[1], ['检查', 'CHECK'])

    def test_duplicate_column_names_and_orders_and_unknown_field_rejected(self):
        for extra, phrase in [({'label_0': 'SAME', 'label_1': 'SAME'}, '列名不能重复'),
                              ({'order_0': '1', 'order_1': '1'}, '顺序不能重复'),
                              ({'fields': ['unknown']}, '选择一个有效的选项')]:
            data = self.data(book_table(self.book), fields=['key', 'name'])
            data.update(extra)
            response = self.client.post(self.url, data)
            self.assertNotIn('attachment;', response.get('Content-Disposition', ''))
            self.assertContains(response, phrase)

    def test_all_formats_enforce_blind_phase_and_final_gate(self):
        for format in ('csv', 'xlsx', 'json', 'tsv'):
            for target in ('all', 'disagreements', 'final'):
                response = self.client.post(self.url, {'target': target, 'round': self.r.id, 'format': format})
                self.assertEqual(response.status_code, 403)
                self.assertNotIn(b'SECRET_OTHER_ANSWER', response.content)
            self.assertEqual(self.client.get(f'/r/{self.r.id}/export/all/?format={format}').status_code, 403)

    def test_mine_is_scoped_to_logged_in_coder_for_every_format(self):
        other = self.own_assignment(self.coder)
        other.note = 'SECRET_OTHER_ANSWER'
        other.save()
        for format in ('csv', 'xlsx', 'json', 'tsv'):
            response = self.client.get(f'/r/{self.r.id}/export/mine/?format={format}')
            self.assertEqual(response.status_code, 200)
            if format == 'xlsx':
                wb = load_workbook(io.BytesIO(response.content), read_only=True)
                raw = str(list(wb.active.values))
                wb.close()
            else:
                raw = response.content.decode('utf-8-sig')
            self.assertNotIn('SECRET_OTHER_ANSWER', raw)
            self.assertIn(self.owner.username, raw)

    def test_formula_like_text_is_escaped_only_in_delimited_files_and_never_executed(self):
        table = ExportTable(['unit_text'], [['=DANGEROUS()'], ['+plus'], ['@name'], ['正常中文']], '虚构')
        csv_response = export_response(table, table.headers, table.headers, 'csv', 'test')
        self.assertIn("'=DANGEROUS()", csv_response.content.decode('utf-8-sig'))
        response = export_response(table, table.headers, ['=HEADING()'], 'xlsx', 'test')
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
        self.assertEqual(workbook.active['A1'].data_type, 's')
        self.assertEqual(workbook.active['A2'].value, '=DANGEROUS()')
        self.assertEqual(workbook.active['A2'].data_type, 's')
        workbook.close()
        response = export_response(table, table.headers, table.headers, 'json', 'test')
        self.assertEqual(json.loads(response.content)['rows'][0]['unit_text'], '=DANGEROUS()')

    def test_final_exports_include_uncertainty_cycle_and_do_not_replace_independent_notes(self):
        self.lock()
        for unit in self.units:
            self.assertEqual(self.decide(unit).status_code, 200)
        response = self.client.post(f'/r/{self.r.id}/manage/', {'revision': self.r.revision, 'action': 'close'})
        self.assertEqual(response.status_code, 302)
        response = self.client.get(f'/r/{self.r.id}/export/final/?format=json')
        payload = json.loads(response.content)
        self.assertEqual(len(payload['rows']), 3)
        self.assertIs(payload['rows'][0]['uncertain'], False)
        self.assertEqual(payload['rows'][0]['cycle'], 1)
        self.assertIn('decision_note', payload['rows'][0])
        self.assertNotIn('coder_note', payload['rows'][0])

    def test_units_include_metadata_and_optional_inactive_only_not_labels(self):
        self.units[0].active = False
        self.units[0].save()
        for include, count in [('', 2), ('yes', 3)]:
            response = self.client.get(self.url, {'target': 'units', 'include_inactive': include})
            self.assertEqual(response.context['row_count'], count)
            self.assertIn('meta_interview', response.context['headers'])
            self.assertNotIn('primary_code', response.context['headers'])

    def test_foreign_project_book_and_round_cannot_be_exported(self):
        self.assertEqual(self.client.get(f'/p/{self.other.id}/exports/').status_code, 404)
        self.assertEqual(self.client.get(self.url, {'book': self.foreign_code.book_id}).status_code, 404)
        self.assertEqual(self.client.get(f'/p/{self.other.id}/exports/', {'target': 'mine', 'round': self.r.id}).status_code, 404)

    def test_unknown_format_and_empty_selection_do_not_download(self):
        data = self.data(book_table(self.book), format='exe')
        response = self.client.post(self.url, data)
        self.assertNotIn('Content-Disposition', response)
        data = self.data(book_table(self.book))
        data['fields'] = []
        response = self.client.post(self.url, data)
        self.assertNotIn('Content-Disposition', response)
        self.assertEqual(self.client.get(f'/r/{self.r.id}/export/mine/?format=exe').status_code, 400)

    def test_large_excel_cell_is_not_silently_truncated(self):
        table = ExportTable(['unit_text'], [['中' * 33000]])
        with self.assertRaisesMessage(ValueError, '不会截断'):
            export_response(table, table.headers, table.headers, 'xlsx', 'test')
        response = export_response(table, table.headers, table.headers, 'json', 'test')
        self.assertEqual(len(json.loads(response.content)['rows'][0]['unit_text']), 33000)

    def test_unsafe_filename_does_not_inject_headers_or_paths(self):
        table = ExportTable(['key'], [['A']])
        response = export_response(table, table.headers, table.headers, 'csv', '../不安全\r\n"name')
        value = response['Content-Disposition']
        self.assertNotIn('\r', value)
        self.assertNotIn('\n', value)
        self.assertNotIn('../', value)

    def test_export_post_requires_csrf(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.owner)
        self.assertEqual(client.post(self.url, self.data(book_table(self.book))).status_code, 403)

    def test_export_row_limit_is_checked_before_building_large_tables(self):
        with patch('coding.file_views.MAX_EXPORT_ROWS', 1):
            response = self.client.get(self.url, {'target': 'mine', 'round': self.r.id})
            self.assertContains(response, '最多导出50000行')
            self.assertNotIn('Content-Disposition', response)
        with patch('coding.exporting.MAX_EXPORT_ROWS', 1):
            self.assertEqual(self.client.get(f'/r/{self.r.id}/export/mine/').status_code, 400)
