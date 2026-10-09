"""Deterministic codebook import. Document prose is never silently treated as rules."""
import csv
import hashlib
import io
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from datetime import timedelta
from xml.etree.ElementTree import ParseError

from defusedxml.ElementTree import fromstring as safe_fromstring
from defusedxml.ElementTree import iterparse as safe_iterparse
from defusedxml.common import DefusedXmlException
from django.db import transaction
from django.utils import timezone

from .importing import MAX_BYTES
from .models import Code, Codebook, CodebookImportDraft, Project
from .services import Conflict, audit

MAX_CODES = 1000
MAX_TEXT = 2 * 1024 * 1024
FIELDS = [
    ('key', '代码标识', True), ('name', '代码名称', True), ('definition', '定义', True),
    ('parent_key', '上级代码标识', False), ('include_when', '适用条件', False),
    ('exclude_when', '排除条件', False), ('positive_example', '正例', False),
    ('negative_example', '反例', False), ('boundary_example', '边界例', False),
    ('coexist_priority', '共现与优先规则', False), ('color', '颜色', False),
]
ALIASES = {
    'key': ['key', 'code', 'code_key', 'code_id', '代码标识', '代码', '编码', '编码标识', '代码编号', '编号'],
    'name': ['name', 'code_name', 'label', '代码名称', '编码名称', '名称'],
    'definition': ['definition', 'description', '定义', '代码定义', '编码定义', '说明'],
    'parent_key': ['parent_key', 'parent_code', 'parent', '上级代码标识', '上级代码', '父代码', '父级代码'],
    'include_when': ['include_when', 'inclusion', 'include', '适用条件', '纳入条件', '使用条件', '什么情况下使用'],
    'exclude_when': ['exclude_when', 'exclusion', 'exclude', '排除条件', '不适用条件', '什么情况下不使用'],
    'positive_example': ['positive_example', 'positive_examples', '正例', '正面例子'],
    'negative_example': ['negative_example', 'negative_examples', '反例', '反面例子'],
    'boundary_example': ['boundary_example', 'boundary_examples', '边界例', '边界案例'],
    'coexist_priority': ['coexist_priority', 'priority', '共现与优先规则', '共现与优先', '共现规则', '优先规则'],
    'color': ['color', 'colour', '颜色'],
}


def normalize_header(value):
    return re.sub(r'[\s\-_：:]+', '', str(value)).casefold()


LOOKUP = {normalize_header(alias): field for field, aliases in ALIASES.items() for alias in aliases}


def suggest_mapping(headers):
    return {field: next((h for h in headers if normalize_header(h) == normalize_header(field)), '')
            or next((h for h in headers if LOOKUP.get(normalize_header(h)) == field), '')
            for field, _, _ in FIELDS}


def text_decode(raw):
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            result = raw.decode(encoding)
            if '\x00' in result:
                raise ValueError('文件含空字符，请另存为UTF-8文本。')
            return result
        except UnicodeError:
            continue
    raise ValueError('无法识别文本编码，请另存为UTF-8。')


def checked_table(headers, rows):
    headers = [str(h).strip() for h in headers]
    if not headers or len(headers) > 100 or any(not h or len(h) > 200 for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('表头不能为空、重复或超过200字；最多100列。')
    if len(rows) > MAX_CODES:
        raise ValueError('每次最多导入1000条代码，请拆分文件。')
    if not rows:
        raise ValueError('没有发现代码记录，请检查文件和工作表。')
    result = []
    total = 0
    for row in rows:
        values = {h: str(row.get(h, '') if row.get(h) is not None else '') for h in headers}
        if any(len(v) > 20000 for v in values.values()):
            raise ValueError('单个字段超过20000字，请检查表格结构。')
        total += sum(len(v) for v in values.values())
        if total > MAX_TEXT:
            raise ValueError('解析文本超过2MB，请拆分编码本。')
        if any(v.strip() for v in values.values()):
            result.append(values)
    if not result:
        raise ValueError('文件中只有空白记录。')
    return headers, result


def csv_table(text, delimiter=','):
    csv.field_size_limit(20000)
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        headers = [h.strip() for h in next(reader, [])]
        rows = []
        for cells in reader:
            if not any(c.strip() for c in cells):
                continue
            if len(cells) > len(headers):
                raise ValueError('数据列比表头多，请检查分隔符。')
            rows.append(dict(zip(headers, cells)))
            if len(rows) > MAX_CODES:
                raise ValueError('每次最多导入1000条代码。')
        return checked_table(headers, rows)
    except csv.Error as error:
        raise ValueError('表格格式错误或单个字段过长。') from error


def inspect_office_zip(raw):
    """Reject entity expansion, duplicate/huge ZIP members and macro-enabled input."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (len(entries) > 500 or sum(e.file_size for e in entries) > 40 * 1024 * 1024
                    or len({e.filename for e in entries}) != len(entries)):
                raise ValueError('文档解压后过大或结构重复，请另存为精简文件。')
            if any('vbaproject' in e.filename.casefold() for e in entries):
                raise ValueError('不导入包含宏的文档。请另存为普通DOCX或XLSX。')
            for entry in entries:
                if entry.filename.lower().endswith(('.xml', '.rels')):
                    for _, element in safe_iterparse(io.BytesIO(archive.read(entry)), events=('end',),
                                                     forbid_dtd=True, forbid_entities=True, forbid_external=True):
                        element.clear()
    except (zipfile.BadZipFile, OSError, RuntimeError, ParseError, DefusedXmlException) as error:
        raise ValueError('Office文件损坏或包含不安全的XML，请另存为普通文档。') from error


def xlsx_table(raw, sheet_name=''):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    inspect_office_zip(raw)
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        names = workbook.sheetnames
        if not names or len(names) > 30:
            raise ValueError('请将所需编码工作表保存到不超过30张表的文件。')
        selected = sheet_name or names[0]
        if not sheet_name:
            for name in names:
                first = next(workbook[name].iter_rows(min_row=1, max_row=1, max_col=100, values_only=True), ())
                if sum(bool(LOOKUP.get(normalize_header(v))) for v in first if v is not None) >= 2:
                    selected = name
                    break
        if selected not in names:
            raise ValueError('请选择文件中实际存在的工作表。')
        sheet = workbook[selected]
        if (sheet.max_row or 0) > MAX_CODES + 1 or (sheet.max_column or 0) > 100:
            raise ValueError('工作表超过1000行代码或100列，请复制实际数据到新工作表。')
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else '' for v in next(iterator, [])]
        rows = []
        for cells in iterator:
            rows.append({h: str(v) if v is not None else '' for h, v in zip(headers, cells)})
            if len(rows) > MAX_CODES:
                raise ValueError('每次最多导入1000条代码。')
        headers, rows = checked_table(headers, rows)
        warnings = ['Excel公式只按原始文字读取，不计算。代码标识建议使用文本单元格，以保留前导零。']
        return headers, rows, names, selected, warnings
    except (zipfile.BadZipFile, KeyError, TypeError, IndexError, OSError, ParseError, InvalidFileException) as error:
        raise ValueError('Excel文件损坏，请另存为XLSX或CSV。') from error
    finally:
        if workbook is not None:
            workbook.close()


def json_table(raw):
    def unique_object(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ValueError('JSON中有重复字段，请删除重复项。')
            result[key] = value
        return result
    try:
        value = json.loads(text_decode(raw), object_pairs_hook=unique_object)
    except (json.JSONDecodeError, RecursionError) as error:
        raise ValueError('JSON格式不正确。') from error
    entries = value.get('codes', value.get('rows')) if isinstance(value, dict) else value
    if not isinstance(entries, list) or not entries or len(entries) > MAX_CODES:
        raise ValueError('JSON需要一个代码对象列表，或包含codes列表；最多1000条。')
    if any(not isinstance(e, dict) for e in entries):
        raise ValueError('JSON中的每条代码需要是字段对象。')
    ids = {str(e['id']): str(e.get('key', '')) for e in entries if 'id' in e}
    rows = []
    for entry in entries:
        entry = dict(entry)
        if ids and 'parent' in entry and 'parent_key' not in entry:
            if entry['parent'] is not None and str(entry['parent']) not in ids:
                raise ValueError('JSON的上级代码编号无法对应到文件中的代码。')
            entry['parent_key'] = ids.get(str(entry['parent']), '')
        if any(isinstance(v, (dict, list)) for v in entry.values()):
            raise ValueError('代码字段应为文字，不支持嵌套字段。')
        rows.append({str(k): '' if v is None else str(v) for k, v in entry.items()})
    headers = list(dict.fromkeys(k for row in rows for k in row))
    return checked_table(headers, rows)


def document_blocks(text):
    if len(text) > MAX_TEXT:
        raise ValueError('文档文字超过2MB，请拆分文件。')
    # Markdown tables have an explicit header and separator; ordinary prose is not guessed.
    lines = text.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    for index in range(len(lines) - 1):
        if '|' in lines[index] and re.fullmatch(r'[\s|:\-]+', lines[index + 1]):
            headers = [c.strip() for c in lines[index].strip().strip('|').split('|')]
            rows = []
            for line in lines[index + 2:]:
                if '|' not in line:
                    break
                cells = [c.strip() for c in line.strip().strip('|').split('|')]
                if len(cells) != len(headers):
                    raise ValueError('Markdown表格列数不一致。单元格中的竖线请改用文字描述。')
                rows.append(dict(zip(headers, cells)))
                if len(rows) > MAX_CODES:
                    raise ValueError('每次最多导入1000条代码。')
            return checked_table(headers, rows), []
    rows, current, last_field, ignored = [], {}, None, 0
    for line in lines:
        match = re.match(r'^\s*(?:[-*]\s+)?([^：:]{1,80})[：:]\s*(.*)$', line)
        field = LOOKUP.get(normalize_header(match[1].strip('* '))) if match else None
        if field:
            if field in current:
                if field in ('key', 'name') and all(current.get(k, '').strip() for k in ('key', 'name', 'definition')):
                    rows.append(current)
                    current = {}
                else:
                    raise ValueError(f'文档中的“{field}”重复且代码边界不清，请按模板分隔每条代码。')
            current[field] = match[2].strip()
            last_field = field
        elif line.strip():
            if last_field and not line.lstrip().startswith('#'):
                current[last_field] += '\n' + line.strip()
            else:
                ignored += 1
        elif current and all(current.get(k, '').strip() for k in ('key', 'name', 'definition')):
            rows.append(current)
            current, last_field = {}, None
        if len(rows) > MAX_CODES:
            raise ValueError('每次最多导入1000条代码。')
    if current:
        rows.append(current)
    if not rows:
        raise ValueError('未识别出编码规则。请使用带表头的表格，或“代码标识：…、代码名称：…、定义：…”的文档模板；普通论文不会自动变成编码本。')
    headers = [f for f, _, _ in FIELDS if any(f in row for row in rows)]
    warnings = ['文档按字段名称解析，不自动推断代码。请逐条核对定义、例子和层级。']
    if ignored:
        warnings.append(f'有{ignored}行标题或说明未作为代码字段导入；原文件仍保留。')
    return checked_table(headers, rows), warnings


def docx_table(raw):
    inspect_office_zip(raw)
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    def paragraph_text(paragraph):
        return ''.join((node.text or '') if node.tag == ns + 't' else '\n' if node.tag == ns + 'br'
                       else '\t' if node.tag == ns + 'tab' else '' for node in paragraph.iter())
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            document = safe_fromstring(archive.read('word/document.xml'), forbid_dtd=True,
                                       forbid_entities=True, forbid_external=True)
        tables, other_tables = [], []
        for table in document.iter(ns + 'tbl'):
            grid = []
            for row in table.findall(ns + 'tr'):
                cells = ['\n'.join(paragraph_text(p)
                                   for p in cell.findall(ns + 'p')) for cell in row.findall(ns + 'tc')]
                grid.append(cells)
            if grid:
                headers = [c.strip() for c in grid[0]]
                if sum(bool(LOOKUP.get(normalize_header(h))) for h in headers) >= 2:
                    tables.append((headers, grid[1:]))
                elif len(headers) >= 3 and all(headers) and len(set(headers)) == len(headers) and len(grid) > 1:
                    other_tables.append((headers, grid[1:]))
        if not tables and other_tables:
            tables = other_tables
        if tables:
            headers = tables[0][0]
            rows = []
            skipped = 0
            for names, grid in tables:
                if names != headers:
                    skipped += 1
                    continue
                for cells in grid:
                    if len(cells) != len(headers):
                        raise ValueError('Word表格含合并单元格或列数不一致，请整理为一行一条代码。')
                    rows.append(dict(zip(headers, cells)))
            warnings = ['Word按有字段表头的表格读取；正文说明不自动变成代码。']
            if skipped:
                warnings.append(f'{skipped}张表的字段不同，未与当前表合并。请分别导入。')
            return checked_table(headers, rows), warnings
        body = document.find(ns + 'body')
        if body is None:
            raise ValueError('Word文档缺少正文。')
        paragraphs = [paragraph_text(p) for p in body.findall(ns + 'p')]
        return document_blocks('\n'.join(paragraphs))
    except (zipfile.BadZipFile, KeyError, ParseError, DefusedXmlException) as error:
        raise ValueError('Word文件损坏，请另存为普通DOCX。') from error


def parse_codebook(filename, raw, sheet_name=''):
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError('请选择非空文件，单个文件最大8MB。')
    suffix = Path(filename).suffix.lower()
    names, selected, warnings = [], '', []
    if suffix == '.xlsx':
        return xlsx_table(raw, sheet_name)
    if suffix in ('.csv', '.tsv'):
        headers, rows = csv_table(text_decode(raw), '\t' if suffix == '.tsv' else ',')
    elif suffix == '.json':
        headers, rows = json_table(raw)
    elif suffix == '.docx':
        (headers, rows), warnings = docx_table(raw)
    elif suffix in ('.txt', '.md'):
        text = text_decode(raw)
        if '\t' in text.split('\n', 1)[0]:
            headers, rows = csv_table(text, '\t')
        else:
            (headers, rows), warnings = document_blocks(text)
    elif suffix == '.pdf':
        # No network, OCR, JavaScript, attachments or external commands inside the PDF.
        # A timed worker keeps problematic PDF parsing off the request thread.
        try:
            worker = subprocess.run([sys.executable, '-X', 'utf8', '-m', 'coding.pdf_extract'], input=raw,
                                    stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=15, check=False,
                                    cwd=Path(__file__).resolve().parent.parent)
            payload = json.loads(worker.stdout.decode('utf-8'))
        except subprocess.TimeoutExpired as error:
            raise ValueError('PDF解析超时。请将规则表格另存为Word或Excel后导入。') from error
        except (OSError, ValueError, UnicodeError) as error:
            raise ValueError('PDF无法安全读取，请将规则另存为Word或Excel。') from error
        if worker.returncode or payload.get('error'):
            raise ValueError(payload.get('error', 'PDF读取失败。'))
        (headers, rows), warnings = document_blocks(payload['text'])
        warnings.append('PDF只提取可复制的文字，不支持扫描图片；多栏、表格和换行可能错位，请认真核对。')
        if payload.get('empty_pages'):
            warnings.append(f'PDF有{payload["empty_pages"]}页没有可提取文字，可能含扫描规则或空白页；这些页未解析，请与原文件核对。')
    else:
        raise ValueError('支持XLSX、CSV、TSV、DOCX、TXT、Markdown、JSON及文字型PDF。旧版XLS或DOC请另存为XLSX或DOCX；扫描PDF先转成文字表格。')
    return headers, rows, names, selected, warnings


def normalized_codes(rows, mapping, existing=()):
    selected = [v for v in mapping.values() if v]
    if len(selected) != len(set(selected)):
        raise ValueError('同一来源列不能同时对应多个编码字段。')
    codes, keys = [], set()
    for position, row in enumerate(rows, 2):
        entry = {field: str(row.get(mapping.get(field), '') or '').strip() for field, _, _ in FIELDS}
        if not any(entry.values()):
            raise ValueError(f'第{position}行有数据，但所选编码字段均为空。请检查列对应，不会静默跳过。')
        for field, label, required in FIELDS:
            value = entry[field]
            if required and not value:
                raise ValueError(f'第{position}行缺少{label}。')
            limit = {'key': 40, 'name': 100, 'parent_key': 40, 'color': 7}.get(field, 5000)
            if len(value) > limit:
                raise ValueError(f'第{position}行{label}超过{limit}字。')
            if any(ord(c) < 32 and c not in '\n\r\t' for c in value):
                raise ValueError(f'第{position}行{label}含无效控制字符。')
        if '\n' in entry['key'] or '\r' in entry['key'] or '\t' in entry['key']:
            raise ValueError(f'第{position}行代码标识不能含换行或制表符。')
        if entry['key'] in keys:
            raise ValueError(f'代码标识“{entry["key"]}”在文件中重复。')
        keys.add(entry['key'])
        if entry['color'] and not re.fullmatch(r'#[0-9a-fA-F]{6}', entry['color']):
            raise ValueError(f'第{position}行颜色需为#加六位十六进制，例如#4372e8。')
        entry['color'] = entry['color'] or '#4372e8'
        codes.append(entry)
    if not codes:
        raise ValueError('没有可导入的代码。')
    parents = {c.key: c.parent.key if c.parent else '' for c in existing}
    overlap = keys & set(parents)
    if overlap:
        raise ValueError('这些标识已存在：' + '、'.join(sorted(overlap)[:8]) + '。不会覆盖旧代码；请选择生成新版本或更改标识。')
    parents.update({c['key']: c['parent_key'] for c in codes})
    for key, parent in parents.items():
        if parent and parent not in parents:
            raise ValueError(f'代码“{key}”的上级“{parent}”不存在；上级请填写代码标识，而不是名称。')
    finished = set()
    for key in parents:
        path, node = set(), key
        while node and node not in finished:
            if node in path:
                raise ValueError(f'代码“{key}”的层级形成循环，请修改上级代码。')
            path.add(node)
            node = parents[node]
        finished.update(path)
    return codes


def draft_available(draft):
    if draft.created_at < timezone.now() - timedelta(hours=24) and not draft.applied_book_id:
        raise ValueError('本次导入预览已超过24小时，请重新上传确认。')


@transaction.atomic
def apply_codebook_import(draft_id, actor, book):
    Project.objects.select_for_update().get(pk=book.project_id)
    locked = Codebook.objects.select_for_update().get(pk=book.pk)
    draft = CodebookImportDraft.objects.select_for_update().get(pk=draft_id, user=actor, book=locked)
    if draft.applied_book_id:
        return draft.applied_book, False
    draft_available(draft)
    preview = draft.preview
    if not preview or preview.get('revision') != locked.revision:
        raise Conflict('预览后编码本已变化，请返回字段对应页面重新预览，不会覆盖他人的修改。')
    settings = preview['settings']
    mode = settings['mode']
    if mode == 'append' and locked.frozen:
        raise Conflict('该编码本已冻结，请重新预览并生成新版本。')
    existing = list(locked.codes.select_related('parent')) if mode == 'append' else []
    codes = normalized_codes(preview['codes'], {f: f for f, _, _ in FIELDS}, existing)
    if mode == 'new':
        version = (locked.project.codebooks.order_by('-version').values_list('version', flat=True).first() or 0) + 1
        target = Codebook.objects.create(project=locked.project, version=version, title=settings['title'],
                                        policy=settings['policy'], change_note=settings['change_note'])
    else:
        target = locked
    objects = {c.key: c for c in existing}
    for entry in codes:
        values = {k: v for k, v in entry.items() if k != 'parent_key'}
        objects[entry['key']] = Code.objects.create(book=target, **values)
    for entry in codes:
        if entry['parent_key']:
            objects[entry['key']].parent = objects[entry['parent_key']]
            objects[entry['key']].save(update_fields=['parent'])
    target.revision += 1
    target.save(update_fields=['revision'])
    draft.applied_book = target
    draft.save(update_fields=['applied_book'])
    audit(locked.project, actor, 'codebook.import', target,
          before={'source_version': locked.version, 'source_revision': preview['revision']},
          after={'version': target.version, 'mode': mode, 'code_count': len(codes), 'filename': draft.filename,
                 'sha256': hashlib.sha256(bytes(draft.raw)).hexdigest(), 'import_id': str(draft.pk),
                 'mapping': preview['mapping'], 'sheet': preview.get('sheet', '')}, reason=settings['change_note'])
    return target, True
