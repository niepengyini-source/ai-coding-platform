"""Portable tabular exports. All formats share one permission-filtered data source."""
import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass
from xml.etree import ElementTree as ET

from django.http import HttpResponse
from django.utils.http import content_disposition_header
from .models import Unit
from .services import csv_safe

MAX_EXPORT_ROWS = 50000
FORMATS = [('xlsx', 'Excel表格（.xlsx）'), ('csv', 'CSV表格（.csv）'),
           ('tsv', '制表符表格（.tsv）'), ('json', 'JSON数据（.json）')]
BASE_LABELS = {
    'project_id': '项目编号', 'round_id': '轮次编号', 'book_version': '编码本版本', 'unit_id': '单元编号',
    'document_id': '原文件编号', 'document_name': '原文件名称', 'record_position': '原记录行号',
    'span_start': '原文起点', 'span_end': '原文终点', 'coder': '编码员', 'unit_text': '编码原文',
    'primary_code': '主代码', 'secondary_code': '次代码', 'uncertain': '是否不确定', 'no_code': '是否无适用代码',
    'coder_note': '编码备注', 'status': '提交状态', 'revision': '保存版本', 'metadata_json': '附加信息JSON',
    'decision_id': '定稿编号', 'decision_note': '定稿依据', 'decided_by': '定稿者', 'cycle': '协商周期',
    'round_kind': '轮次类型', 'disagreement_types': '分歧类型', 'resolved': '是否已解决',
    'independent_labels_json': '独立判断JSON', 'active': '单元是否有效', 'unit_note': '单元调整说明',
    'key': '代码标识', 'name': '代码名称', 'definition': '定义', 'parent_key': '上级代码标识',
    'include_when': '适用条件', 'exclude_when': '排除条件', 'positive_example': '正例',
    'negative_example': '反例', 'boundary_example': '边界例', 'coexist_priority': '共现与优先规则', 'color': '颜色',
}


@dataclass
class ExportTable:
    headers: list
    rows: list
    title: str = ''
    policy: str = ''


def metadata_columns(units):
    return sorted({str(k) for unit in units for k in unit.record.metadata})


def annotation_table(r, queryset):
    codes = dict(r.book.codes.values_list('id', 'key'))
    entries = list(queryset.select_related('unit__record', 'coder').order_by('unit_id', 'coder_id'))
    # Scope metadata to the same visible assignments; never reveal other members' records.
    metadata_keys = metadata_columns([a.unit for a in entries])
    headers = ['project_id', 'round_id', 'book_version', 'unit_id', 'document_id', 'record_position',
               'span_start', 'span_end', 'coder', 'unit_text', 'primary_code', 'secondary_code',
               'uncertain', 'no_code', 'coder_note', 'status', 'revision']
    rows = []
    for a in entries:
        rows.append([r.project_id, r.id, r.book.version, a.unit_id, a.unit.record.document_id, a.unit.record.position,
                     a.unit.start, a.unit.end, a.coder.username, a.unit.text, codes.get(a.primary_id, ''),
                     '|'.join(codes.get(c, '') for c in a.secondary), a.uncertain, a.no_code, a.note, a.status, a.revision]
                    + [a.unit.record.metadata.get(k, '') for k in metadata_keys]
                    + [json.dumps(a.unit.record.metadata, ensure_ascii=False)])
    return ExportTable(headers + ['meta_' + k for k in metadata_keys] + ['metadata_json'], rows, r.name)


def final_table(r):
    codes = dict(r.book.codes.values_list('id', 'key'))
    entries = list(r.decisions.filter(cycle=r.cycle).select_related('unit__record', 'author').order_by('unit_id'))
    metadata_keys = metadata_columns([a.unit for a in entries])
    headers = ['project_id', 'round_id', 'book_version', 'unit_id', 'document_id', 'record_position',
               'span_start', 'span_end', 'unit_text', 'primary_code', 'secondary_code', 'no_code', 'decision_id',
               'decision_note', 'decided_by', 'revision', 'uncertain', 'cycle']
    rows = []
    for a in entries:
        rows.append([r.project_id, r.id, r.book.version, a.unit_id, a.unit.record.document_id, a.unit.record.position,
                     a.unit.start, a.unit.end, a.unit.text, codes.get(a.primary_id, ''),
                     '|'.join(codes.get(c, '') for c in a.secondary), a.no_code, a.id, a.note, a.author.username,
                     a.revision, a.uncertain, a.cycle]
                    + [a.unit.record.metadata.get(k, '') for k in metadata_keys]
                    + [json.dumps(a.unit.record.metadata, ensure_ascii=False)])
    return ExportTable(headers + ['meta_' + k for k in metadata_keys] + ['metadata_json'], rows, r.name)


def disagreement_table(r):
    from .reviewing import disagreement_rows
    codes = dict(r.book.codes.values_list('id', 'key'))
    rows = []
    for item in disagreement_rows(r):
        if not item['needs_attention']:
            continue
        independent = [{'coder': a.coder.username, 'primary_code': codes.get(a.primary_id, ''),
                        'secondary_code': [codes.get(c, '') for c in a.secondary], 'uncertain': a.uncertain,
                        'no_code': a.no_code, 'status': a.status, 'coder_note': a.note, 'revision': a.revision}
                       for a in item['entries']]
        unit, decision = item['unit'], item['decision']
        rows.append([r.id, r.kind, r.cycle, r.book.version, unit.id, unit.text, '|'.join(item['types']),
                     item['resolved'], decision.id if decision else '', json.dumps(independent, ensure_ascii=False),
                     json.dumps(unit.record.metadata, ensure_ascii=False)])
    return ExportTable(['round_id', 'round_kind', 'cycle', 'book_version', 'unit_id', 'unit_text', 'disagreement_types',
                        'resolved', 'decision_id', 'independent_labels_json', 'metadata_json'], rows, r.name)


def unit_table(project, include_inactive=False):
    queryset = Unit.objects.filter(record__document__project=project).select_related('record__document')
    if not include_inactive:
        queryset = queryset.filter(active=True)
    units = list(queryset.order_by('record__document_id', 'record__position', 'start', 'id')[:MAX_EXPORT_ROWS + 1])
    if len(units) > MAX_EXPORT_ROWS:
        raise ValueError('一次最多导出50000行，请按轮次或编码员分别导出。')
    metadata_keys = metadata_columns(units)
    headers = ['project_id', 'unit_id', 'document_id', 'document_name', 'record_position', 'span_start', 'span_end',
               'unit_text', 'active', 'unit_note'] + ['meta_' + k for k in metadata_keys] + ['metadata_json']
    rows = [[project.id, u.id, u.record.document_id, u.record.document.name, u.record.position, u.start, u.end,
             u.text, u.active, u.note] + [u.record.metadata.get(k, '') for k in metadata_keys]
            + [json.dumps(u.record.metadata, ensure_ascii=False)] for u in units]
    return ExportTable(headers, rows, project.name)


def book_table(book):
    from .codebook_files import FIELDS
    headers = [f for f, _, _ in FIELDS]
    rows = [[(c.parent.key if c.parent else '') if f == 'parent_key' else getattr(c, f)
             for f in headers] for c in book.codes.select_related('parent').order_by('key')]
    return ExportTable(headers, rows, book.title, book.policy)


def xlsx_bytes(headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter
    if any(isinstance(value, str) and len(value) > 32767 for row in rows for value in row):
        raise ValueError('某字段超过Excel单元格32767字限制，请改用JSON或CSV，平台不会截断内容。')
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('编码数据')
    sheet.freeze_panes = 'A2'
    for index, header in enumerate(headers, 1):
        width = 44 if any(token in header.casefold() for token in ('text', 'note', 'definition', 'example', 'json')) else 24
        sheet.column_dimensions[get_column_letter(index)].width = width
    heading = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.data_type = 's'
        cell.font = Font(name='宋体', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='34496B')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        heading.append(cell)
    sheet.append(heading)
    for row in rows:
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'  # Do not turn research text or a custom heading into a formula.
            cell.font = Font(name='宋体', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def docx_bytes(headers, rows, title):
    """Small, standards-based OOXML with text only; no macros, fields or external relationships."""
    w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    ET.register_namespace('w', w)
    document = ET.Element(f'{{{w}}}document')
    body = ET.SubElement(document, f'{{{w}}}body')
    def paragraph(text='', heading=False):
        p = ET.SubElement(body, f'{{{w}}}p')
        if heading:
            prop = ET.SubElement(p, f'{{{w}}}pPr')
            ET.SubElement(prop, f'{{{w}}}pStyle', {f'{{{w}}}val': 'Title'})
        run = ET.SubElement(p, f'{{{w}}}r')
        for index, line in enumerate(str(text).split('\n')):
            if index:
                ET.SubElement(run, f'{{{w}}}br')
            node = ET.SubElement(run, f'{{{w}}}t', {'{http://www.w3.org/XML/1998/namespace}space': 'preserve'})
            node.text = line
    paragraph(title, heading=True)
    paragraph()
    for row in rows:
        for header, value in zip(headers, row):
            paragraph(f'{header}: {value}')
        paragraph()
    section = ET.SubElement(body, f'{{{w}}}sectPr')
    ET.SubElement(section, f'{{{w}}}pgSz', {f'{{{w}}}w': '11906', f'{{{w}}}h': '16838'})
    ET.SubElement(section, f'{{{w}}}pgMar', {f'{{{w}}}top': '1440', f'{{{w}}}bottom': '1440', f'{{{w}}}left': '1440', f'{{{w}}}right': '1440'})
    content_types = '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>'
    root_rels = '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    doc_rels = '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'
    styles = f'''<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="{w}"><w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="宋体"/><w:sz w:val="24"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="120" w:line="360" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults><w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:pPr><w:keepNext/></w:pPr><w:rPr><w:rFonts w:eastAsia="黑体"/><w:b/><w:color w:val="000000"/><w:sz w:val="32"/></w:rPr></w:style></w:styles>'''
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in {'[Content_Types].xml': content_types, '_rels/.rels': root_rels,
                           'word/_rels/document.xml.rels': doc_rels, 'word/styles.xml': styles,
                           'word/document.xml': ET.tostring(document, encoding='utf-8', xml_declaration=True)}.items():
            archive.writestr(name, data)
    return output.getvalue()


def export_response(table, fields, labels, file_format, filename, *, encoding='utf-8-sig', delimiter=',', is_book=False):
    if len(table.rows) > MAX_EXPORT_ROWS:
        raise ValueError('一次最多导出50000行，请按轮次或编码员分别导出。')
    if (any(field not in table.headers for field in fields) or not fields or len(fields) > 300
            or len(set(fields)) != len(fields) or len(labels) != len(fields) or len(set(labels)) != len(labels)):
        raise ValueError('请选择有效字段，最多导出300列。')
    indexes = [table.headers.index(field) for field in fields]
    rows = [[row[i] for i in indexes] for row in table.rows]
    for row in rows:
        if any(isinstance(v, str) and any(ord(c) < 32 and c not in '\n\r\t' for c in v) for v in row):
            if file_format in ('xlsx', 'docx'):
                raise ValueError('数据含Office不支持的控制字符，请选择JSON或CSV；平台不会删除原文。')
    if file_format in ('csv', 'tsv'):
        stream = io.StringIO(newline='')
        writer = csv.writer(stream, delimiter='\t' if file_format == 'tsv' else delimiter)
        writer.writerow([csv_safe(label) for label in labels])
        writer.writerows([[csv_safe(value) for value in row] for row in rows])
        try:
            data = stream.getvalue().encode(encoding)
        except UnicodeEncodeError as error:
            raise ValueError('所选编码无法保存某些字符，请选择UTF-8，平台不会用问号替换原文。') from error
        content_type = 'text/tab-separated-values' if file_format == 'tsv' else 'text/csv'
        content_type += '; charset=' + ('utf-8' if encoding == 'utf-8-sig' else encoding)
    elif file_format == 'xlsx':
        data = xlsx_bytes(labels, rows)
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif file_format == 'json':
        payload = {'format_version': 1, 'title': table.title,
                   'columns': [{'source': field, 'name': label} for field, label in zip(fields, labels)],
                   'row_count': len(rows), 'codes' if is_book else 'rows': [dict(zip(labels, row)) for row in rows]}
        if is_book:
            payload['policy'] = table.policy
        data = json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8')
        content_type = 'application/json; charset=utf-8'
    elif file_format == 'docx' and is_book:
        data = docx_bytes(labels, rows, table.title)
        content_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        raise ValueError('不支持这个导出格式；Word仅用于编码本。')
    filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', filename).strip(' .')[:100] or '编码数据'
    response = HttpResponse(data, content_type=content_type)
    response['Content-Disposition'] = content_disposition_header(True, filename + '.' + file_format)
    response['Cache-Control'] = 'no-store'
    return response
