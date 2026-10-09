import csv
import io
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError
from defusedxml.ElementTree import iterparse as safe_iterparse
from defusedxml.common import DefusedXmlException

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 10000


def parse_upload(filename, raw):
    if len(raw) > MAX_BYTES:
        raise ValueError('单个文件最大8MB，请按材料批次拆分。')
    suffix = Path(filename).suffix.lower()
    if suffix in ('.csv', '.txt'):
        text = None
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                continue
        if text is None:
            raise ValueError('无法识别文本编码，请另存为UTF-8。')
        if '\x00' in text:
            raise ValueError('材料中含有无效空字符。')
        if suffix == '.txt':
            lines = text.replace('\r\n', '\n').replace('\r', '\n').split('\n')
            rows = [{'text': line} for line in lines]
            headers = ['text']
        else:
            csv.field_size_limit(200000)
            rows = []
            try:
                reader = csv.DictReader(io.StringIO(text), strict=True)
                headers = reader.fieldnames or []
                for row in reader:
                    if None in row:
                        raise ValueError('存在比表头更多的列，请检查CSV格式。')
                    rows.append({k: v or '' for k, v in row.items()})
                    if len(rows) > MAX_ROWS:
                        break
            except csv.Error as error:
                raise ValueError('CSV格式错误或单个单元格过长。') from error
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        wb = None
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 200 or sum(i.file_size for i in entries) > 40 * 1024 * 1024:
                    raise ValueError('Excel解压后过大，请拆分文件。')
                # Validate every XML part before openpyxl, even if lxml is installed later.
                # Never resolve entities, DTDs or external documents from an upload.
                for entry in entries:
                    if entry.filename.lower().endswith(('.xml', '.rels')):
                        for _, element in safe_iterparse(io.BytesIO(archive.read(entry)), events=('end',),
                                                         forbid_dtd=True, forbid_entities=True, forbid_external=True):
                            element.clear()
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            if len(wb.worksheets) != 1:
                raise ValueError('请将需要导入的工作表单独保存为一个文件。')
            sheet = wb.active
            if sheet.max_column > 100 or sheet.max_row > MAX_ROWS + 1:
                raise ValueError('工作表超过100列或10000条记录，请只复制实际材料到新工作表。')
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, ())
            headers = [str(x) if x is not None else '' for x in first]
            rows = []
            for values in iterator:
                rows.append({k: str(v) if v is not None else '' for k, v in zip(headers, values)})
                if len(rows) > MAX_ROWS:
                    break
        except (zipfile.BadZipFile, KeyError, OSError, StopIteration, TypeError, IndexError,
                ParseError, DefusedXmlException, InvalidFileException) as error:
            raise ValueError('Excel文件损坏或不支持，请另存为.xlsx或CSV。') from error
        finally:
            if wb is not None:
                wb.close()
    else:
        raise ValueError('首版支持CSV、XLSX和TXT。旧版XLS请保留原件并另存为XLSX。')
    if not headers or any(not h.strip() or len(h) > 200 for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('表头不能为空、重复或超过200个字符。')
    if len(headers) > 100:
        raise ValueError('最多支持100列。')
    if len(rows) > MAX_ROWS:
        raise ValueError('单批最多10000行，请拆分导入。')
    if any(len(str(value)) > 100000 for row in rows for value in row.values()):
        raise ValueError('单元格过长，请拆分材料。')
    return headers, rows
